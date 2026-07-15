import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from xperimental_data_conv import (
    FlapjackPlateExporter,
    SampleDesignMetadata,
    SynBioHubSampleDesignResolver,
    TemplateSample,
)


FIXTURES = Path(__file__).parent / "test_files"


class FakeObject:
    def __init__(
        self,
        display_id,
        *,
        roles=(),
        properties=None,
        modules=(),
        functional_components=(),
    ):
        self.displayId = display_id
        self.roles = roles
        self.properties = properties or {}
        self.modules = modules
        self.functionalComponents = functional_components


class FakeReference:
    def __init__(self, definition, display_id):
        self.definition = definition
        self.displayId = display_id


class FakeSynBioHubClient:
    def __init__(self, objects):
        self.objects = objects

    def get_object(self, uri):
        return self.objects[uri]


class StaticResolver:
    def __init__(self):
        self.resolved_uris = []

    def resolve(self, sample):
        self.resolved_uris.append(sample.sample_design_uri)
        return SampleDesignMetadata(
            medium="m9",
            strain="Ecolisc2",
            plasmids=("pIJAI477", "pIJAI478"),
            chemicals={"Arabinose": 0.004},
        )


class TestSynBioHubSampleDesignResolver(unittest.TestCase):
    def test_resolves_sbol_modules_by_roles_and_extended_properties(self):
        sample_uri = "https://example.org/sample/1"
        medium_uri = "https://example.org/m9/1"
        chemical_module_uri = "https://example.org/arabinose_004/1"
        chemical_uri = "https://example.org/arabinose/1"
        strain_uri = "https://example.org/strain/1"
        plasmid_uris = ["https://example.org/p{0}/1".format(index) for index in range(1, 5)]

        objects = {
            sample_uri: FakeObject(
                "sample_design",
                modules=(
                    FakeReference(medium_uri, "m9"),
                    FakeReference(chemical_module_uri, "arabinose_004"),
                    FakeReference(strain_uri, "Ecolisc2"),
                ),
            ),
            medium_uri: FakeObject(
                "m9", roles=("http://identifiers.org/ncit/NCIT:C48164",)
            ),
            chemical_module_uri: FakeObject(
                "arabinose_004",
                properties={
                    "https://wiki.synbiohub.org/wiki/Terms/Flapjack#concentration": [
                        "0.004"
                    ]
                },
                functional_components=(FakeReference(chemical_uri, "Arabinose"),),
            ),
            chemical_uri: FakeObject("Arabinose"),
            strain_uri: FakeObject(
                "Ecolisc2",
                roles=("https://identifiers.org/obo/ncit:C14419",),
                functional_components=tuple(
                    FakeReference(uri, "p{0}".format(index))
                    for index, uri in enumerate(plasmid_uris, start=1)
                ),
            ),
        }
        objects.update(
            {
                uri: FakeObject("pIJAI{0}".format(476 + index))
                for index, uri in enumerate(plasmid_uris, start=1)
            }
        )

        resolver = SynBioHubSampleDesignResolver(FakeSynBioHubClient(objects))
        metadata = resolver.resolve(
            TemplateSample("sample1", "A", 1, "sample_design", sample_uri)
        )

        self.assertEqual("m9", metadata.medium)
        self.assertEqual("Ecolisc2", metadata.strain)
        self.assertEqual(
            ("pIJAI477", "pIJAI478", "pIJAI479", "pIJAI480"),
            metadata.plasmids,
        )
        self.assertEqual({"Arabinose": 0.004}, metadata.chemicals)


class TestFlapjackPlateExporter(unittest.TestCase):
    def test_creates_flapjack_workbook_from_align_tf_files(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "flapjack_input.xlsx"
            resolver = StaticResolver()
            result = FlapjackPlateExporter(
                FIXTURES / "Align_TF_Study.xlsm", resolver
            ).export(
                FIXTURES / "2025-07-02_Align-TF-Cytom-12-plasmids_growth-plate_1-1_Neo.txt",
                output,
            )

            self.assertEqual(output, result)
            self.assertTrue(resolver.resolved_uris)
            workbook = load_workbook(output, data_only=True, read_only=True)
            try:
                self.assertEqual(
                    ["Data", "Media", "Strains", "DNA", "Chemicals"],
                    workbook.sheetnames,
                )
                self.assertEqual("Software Version", workbook["Data"]["A2"].value)
                self.assertEqual("3.11.19", workbook["Data"]["B2"].value)
                self.assertEqual("m9", workbook["Media"]["B2"].value)
                self.assertEqual("Ecolisc2", workbook["Strains"]["B2"].value)
                self.assertEqual("DNA 1", workbook["DNA"]["A1"].value)
                self.assertEqual("pIJAI477", workbook["DNA"]["B2"].value)
                self.assertEqual("DNA 2", workbook["DNA"]["A11"].value)
                self.assertEqual("pIJAI478", workbook["DNA"]["B12"].value)
                self.assertEqual("Arabinose", workbook["Chemicals"]["A1"].value)
                self.assertEqual(0.004, workbook["Chemicals"]["B2"].value)
            finally:
                workbook.close()
