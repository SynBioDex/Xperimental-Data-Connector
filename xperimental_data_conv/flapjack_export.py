"""Create Flapjack plate-reader workbooks from XDC and SynBioHub SBOL data."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, Mapping, Optional, Sequence, Tuple, Union

from openpyxl import Workbook, load_workbook


_PLATE_ROWS = tuple("ABCDEFGH")
_PLATE_COLUMNS = tuple(range(1, 13))
_ID_NORMALIZER = re.compile(r"[^a-z0-9]+")
_INTEGER = re.compile(r"[+-]?\d+")
_NUMBER = re.compile(r"[+-]?(?:\d+\.\d*|\d*\.\d+|\d+)(?:[eE][+-]?\d+)?")
_MEDIUM_ROLE = "C48164"
_STRAIN_ROLE = "C14419"
_CONCENTRATION_PROPERTY = "Flapjack#concentration"


def _normalise_identifier(value: str) -> str:
    return _ID_NORMALIZER.sub("_", value.lower()).strip("_")


@dataclass(frozen=True)
class TemplateSample:
    """A sample position and its SynBioHub sample-design URI."""

    sample_id: str
    row: str
    column: int
    sample_design: str
    sample_design_uri: str


@dataclass(frozen=True)
class SampleDesignMetadata:
    """Flapjack metadata resolved from an SBOL sample design."""

    medium: Optional[str]
    strain: Optional[str]
    plasmids: Tuple[str, ...]
    chemicals: Mapping[str, float]


class SynBioHubClient:
    """Authenticated pySBOL2 access to SynBioHub objects.

    Provide either ``token`` or both ``username`` and ``password``. Omitting
    all credentials allows access to publicly readable objects.
    """

    def __init__(
        self,
        server_url: str = "https://synbiohub.org",
        *,
        username: Optional[str] = None,
        password: Optional[str] = None,
        token: Optional[str] = None,
    ):
        if token and (username or password):
            raise ValueError("Use either token authentication or username/password authentication")
        if username and not password:
            raise ValueError("A password is required when a SynBioHub username is supplied")

        self.server_url = server_url.rstrip("/")
        self.username = username
        self.password = password
        self.token = token
        self._part_shop = None
        self._objects: Dict[str, object] = {}

    def get_object(self, uri: str):
        """Pull an SBOL object recursively and return its top-level object."""

        if uri not in self._objects:
            import sbol2

            document = sbol2.Document()
            self._shop().pull(uri, document, recursive=True)
            self._objects[uri] = document.getTopLevel(uri)
        return self._objects[uri]

    def _shop(self):
        if self._part_shop is None:
            import sbol2

            self._part_shop = sbol2.PartShop(self.server_url)
            if self.token:
                # pySBOL2 sends this value as the X-authorization header.
                self._part_shop.key = self.token
            elif self.username:
                self._part_shop.login(self.username, self.password)
        return self._part_shop


class SynBioHubSampleDesignResolver:
    """Resolve media, chemicals, strain, and plasmids from SBOL sample designs."""

    def __init__(self, client: SynBioHubClient):
        self.client = client

    def resolve(self, sample: TemplateSample) -> SampleDesignMetadata:
        sample_design = self.client.get_object(sample.sample_design_uri)
        medium = None
        strain = None
        plasmids = []
        chemicals = {}

        for module in sample_design.modules:
            definition = self.client.get_object(str(module.definition))
            roles = tuple(str(role) for role in definition.roles)

            if self._has_role(roles, _MEDIUM_ROLE):
                medium = self._display_name(definition)
            elif self._has_role(roles, _STRAIN_ROLE):
                strain = self._display_name(definition)
                plasmids.extend(self._plasmids(definition))
            elif self._concentration(definition) is not None:
                chemicals.update(self._chemicals(definition))
            else:
                raise ValueError(
                    "Cannot classify sample-design module '{0}' ({1})"
                    .format(module.displayId, module.definition)
                )

        return SampleDesignMetadata(
            medium=medium,
            strain=strain,
            plasmids=tuple(plasmids),
            chemicals=chemicals,
        )

    @staticmethod
    def _has_role(roles: Sequence[str], ncit_identifier: str) -> bool:
        return any(role.rstrip("/").endswith(ncit_identifier) for role in roles)

    @staticmethod
    def _display_name(sbol_object) -> str:
        return str(sbol_object.displayId)

    def _plasmids(self, strain_definition) -> Iterable[str]:
        for component in strain_definition.functionalComponents:
            plasmid = self.client.get_object(str(component.definition))
            yield self._display_name(plasmid)

    def _chemicals(self, chemical_definition) -> Mapping[str, float]:
        concentration = self._concentration(chemical_definition)
        chemicals = {}
        for component in chemical_definition.functionalComponents:
            chemical = self.client.get_object(str(component.definition))
            chemicals[self._display_name(chemical)] = concentration
        if not chemicals:
            raise ValueError(
                "Chemical module '{0}' has a concentration but no functional components"
                .format(chemical_definition.displayId)
            )
        return chemicals

    @staticmethod
    def _concentration(sbol_object) -> Optional[float]:
        for property_uri, values in sbol_object.properties.items():
            if str(property_uri).endswith(_CONCENTRATION_PROPERTY):
                return float(values[0])
        return None


class FlapjackPlateExporter:
    """Create a Flapjack XLSX workbook from an XDC template and reader export."""

    def __init__(
        self,
        xdc_template: Union[str, Path],
        sample_design_resolver: SynBioHubSampleDesignResolver,
    ):
        self.xdc_template = Path(xdc_template)
        self.sample_design_resolver = sample_design_resolver

    def export(
        self,
        reader_export: Union[str, Path],
        output_path: Union[str, Path],
        *,
        assay_id: Optional[str] = None,
    ) -> Path:
        """Write a Flapjack workbook and return its path."""

        reader_export = Path(reader_export)
        output_path = Path(output_path)
        samples = self._samples_for_assay(reader_export, assay_id)
        metadata = {}
        for sample in samples:
            try:
                metadata[sample.sample_id] = self.sample_design_resolver.resolve(sample)
            except Exception as error:
                raise RuntimeError(
                    "Could not resolve sample '{0}' from SynBioHub URI: {1}"
                    .format(sample.sample_id, sample.sample_design_uri)
                ) from error

        workbook = Workbook()
        data_sheet = workbook.active
        data_sheet.title = "Data"
        self._write_reader_export(data_sheet, reader_export)
        self._write_metadata_sheets(workbook, samples, metadata)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
        return output_path

    def _samples_for_assay(
        self, reader_export: Path, assay_id: Optional[str]
    ) -> Iterable[TemplateSample]:
        workbook = load_workbook(self.xdc_template, data_only=True, read_only=True)
        try:
            sample_sheet = self._find_sheet(workbook, "sample")
            sample_design_uris = self._sample_design_uris(workbook)
            headers = self._headers(sample_sheet)
            required = ("sample id", "row", "column", "assay id", "sample design")
            missing = [header for header in required if header not in headers]
            if missing:
                raise ValueError(
                    "The XDC sample sheet is missing required columns: "
                    + ", ".join(missing)
                )

            records = []
            for row in sample_sheet.iter_rows(min_row=2, values_only=True):
                sample_id = row[headers["sample id"]]
                row_name = row[headers["row"]]
                column = row[headers["column"]]
                current_assay = row[headers["assay id"]]
                design = row[headers["sample design"]]
                if not all((sample_id, row_name, column, current_assay, design)):
                    continue
                design_name = str(design)
                if design_name not in sample_design_uris:
                    raise ValueError(
                        "Sample design '{0}' is missing from SBH_sampledesigns_collection"
                        .format(design_name)
                    )
                sample = TemplateSample(
                    sample_id=str(sample_id),
                    row=str(row_name).upper(),
                    column=int(column),
                    sample_design=design_name,
                    sample_design_uri=sample_design_uris[design_name],
                )
                records.append((str(current_assay), sample))
        finally:
            workbook.close()

        selected_assay = assay_id or self._infer_assay_id(reader_export, records)
        samples = [sample for current_assay, sample in records if current_assay == selected_assay]
        if not samples:
            raise ValueError("No samples found for assay ID: {0}".format(selected_assay))
        invalid = [
            sample
            for sample in samples
            if sample.row not in _PLATE_ROWS or sample.column not in _PLATE_COLUMNS
        ]
        if invalid:
            raise ValueError("The XDC template has samples outside the supported 96-well plate")
        return samples

    @classmethod
    def _sample_design_uris(cls, workbook) -> Mapping[str, str]:
        sheet = cls._find_sheet(workbook, "SBH_sampledesigns_collection")
        headers = cls._headers(sheet)
        required = ("name", "uri")
        missing = [header for header in required if header not in headers]
        if missing:
            raise ValueError(
                "SBH_sampledesigns_collection is missing required columns: "
                + ", ".join(missing)
            )
        return {
            str(row[headers["name"]]): str(row[headers["uri"]])
            for row in sheet.iter_rows(min_row=2, values_only=True)
            if row[headers["name"]] and row[headers["uri"]]
        }

    @staticmethod
    def _headers(sheet) -> Mapping[str, int]:
        return {
            str(cell.value).strip().lower(): index
            for index, cell in enumerate(next(sheet.iter_rows()))
            if cell.value is not None
        }

    @staticmethod
    def _find_sheet(workbook, name: str):
        for sheet_name in workbook.sheetnames:
            if sheet_name.lower() == name.lower():
                return workbook[sheet_name]
        raise ValueError("The XDC template does not contain a '{0}' worksheet".format(name))

    @staticmethod
    def _infer_assay_id(reader_export: Path, records) -> str:
        file_id = _normalise_identifier(reader_export.stem)
        assay_ids = sorted({assay_id for assay_id, _ in records}, key=len, reverse=True)
        matches = [
            assay_id
            for assay_id in assay_ids
            if _normalise_identifier(assay_id) in file_id
        ]
        if len(matches) == 1:
            return matches[0]
        if not matches:
            raise ValueError(
                "Could not infer an assay ID from '{0}'. Pass assay_id explicitly."
                .format(reader_export.name)
            )
        raise ValueError("Reader file name matches multiple assays; pass assay_id explicitly.")

    @staticmethod
    def _write_reader_export(sheet, reader_export: Path) -> None:
        # Flapjack's reference workbooks reserve the first row in Data.
        sheet.append([None])
        with reader_export.open(
            "r", encoding="utf-8-sig", errors="replace", newline=""
        ) as source:
            for row in csv.reader(source):
                sheet.append([FlapjackPlateExporter._excel_value(value) for value in row])

    @staticmethod
    def _excel_value(value: str):
        value = value.strip()
        if not value:
            return None
        if _INTEGER.fullmatch(value):
            return int(value)
        if _NUMBER.fullmatch(value):
            return float(value)
        return value

    def _write_metadata_sheets(self, workbook, samples, metadata) -> None:
        media = workbook.create_sheet("Media")
        self._write_grid(
            media,
            "Media",
            samples,
            {sample.sample_id: metadata[sample.sample_id].medium for sample in samples},
        )

        strains = workbook.create_sheet("Strains")
        self._write_grid(
            strains,
            "Strains",
            samples,
            {sample.sample_id: metadata[sample.sample_id].strain for sample in samples},
        )

        dna = workbook.create_sheet("DNA")
        plasmid_count = max(len(value.plasmids) for value in metadata.values())
        for index in range(plasmid_count):
            self._write_grid(
                dna,
                "DNA {0}".format(index + 1),
                samples,
                {
                    sample.sample_id: self._nth_or_none(
                        metadata[sample.sample_id].plasmids, index
                    )
                    for sample in samples
                },
                start_row=index * 10 + 1,
            )

        chemicals = workbook.create_sheet("Chemicals")
        chemical_names = sorted(
            {
                chemical
                for value in metadata.values()
                for chemical in value.chemicals
            }
        )
        if chemical_names:
            for index, chemical_name in enumerate(chemical_names):
                self._write_grid(
                    chemicals,
                    chemical_name,
                    samples,
                    {
                        sample.sample_id: metadata[sample.sample_id].chemicals.get(chemical_name)
                        for sample in samples
                    },
                    start_row=index * 10 + 1,
                )
        else:
            self._write_grid(chemicals, "Chemicals", samples, {}, start_row=1)

    @staticmethod
    def _nth_or_none(values: Sequence[str], index: int) -> Optional[str]:
        return values[index] if index < len(values) else None

    @staticmethod
    def _write_grid(sheet, title: str, samples, values, *, start_row: int = 1) -> None:
        sheet.cell(start_row, 1, title)
        for column in _PLATE_COLUMNS:
            sheet.cell(start_row, column + 1, column)
        for row_index, row_name in enumerate(_PLATE_ROWS, start=start_row + 1):
            sheet.cell(row_index, 1, row_name)

        for sample in samples:
            value = values.get(sample.sample_id)
            sheet.cell(
                start_row + _PLATE_ROWS.index(sample.row) + 1,
                sample.column + 1,
                "None" if value is None else value,
            )


def create_flapjack_input(
    xdc_template: Union[str, Path],
    reader_export: Union[str, Path],
    output_path: Union[str, Path],
    *,
    assay_id: Optional[str] = None,
    synbiohub_url: str = "https://synbiohub.org",
    username: Optional[str] = None,
    password: Optional[str] = None,
    token: Optional[str] = None,
) -> Path:
    """Create a Flapjack workbook using the XDC template's SBOL design URIs."""

    client = SynBioHubClient(
        synbiohub_url,
        username=username,
        password=password,
        token=token,
    )
    resolver = SynBioHubSampleDesignResolver(client)
    return FlapjackPlateExporter(xdc_template, resolver).export(
        reader_export,
        output_path,
        assay_id=assay_id,
    )
